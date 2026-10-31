# -*- coding: utf-8 -*-
"""把静态页 PNG 组装成 16:9 PPTX，便于预览/投屏/录制。"""

from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[2]
SRC = ROOT / "demo_video" / "静态页"
OUT = SRC / "静态页总览.pptx"

pngs = sorted(SRC.glob("*.png"))
if not pngs:
    raise SystemExit("没有静态页 PNG")

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank = prs.slide_layouts[6]

for png in pngs:
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(str(png), 0, 0, width=prs.slide_width, height=prs.slide_height)
    slide.notes_slide.notes_text_frame.text = png.stem

prs.save(OUT)
print(f"saved {OUT}  slides={len(prs.slides)}")
